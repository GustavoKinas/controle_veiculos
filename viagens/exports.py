"""
Geração de planilhas Excel (.xlsx) do módulo de viagens.

Mantido fora das views pelo mesmo motivo de services.py: reuso (link direto
após concluir o fechamento e botão na listagem) e testabilidade isolada.
"""

import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import Fechamento

NEGRITO = Font(bold=True)


def _autofit(ws) -> None:
    for coluna_celulas in ws.columns:
        maior = max(
            (len(str(c.value)) for c in coluna_celulas if c.value is not None),
            default=0,
        )
        letra = get_column_letter(coluna_celulas[0].column)
        ws.column_dimensions[letra].width = min(maior + 2, 40)


def exportar_fechamento_excel(fechamento: Fechamento) -> HttpResponse:
    """
    Gera o .xlsx de um fechamento: aba "Resumo do Rateio" (por Centro de
    Custo) e aba "Viagens" (detalhamento de cada viagem incluída).

    Usa o modo padrão do openpyxl (não write_only) para poder reler as
    células no autofit de colunas — plenamente adequado ao volume de um
    fechamento por período (dezenas a poucas centenas de viagens). Para
    exportações muito maiores (dezenas de milhares de linhas), o modo
    write_only seria a escolha certa, abrindo mão do autofit.
    """
    wb = Workbook()

    resumo = wb.active
    resumo.title = "Resumo do Rateio"
    resumo.append(["Centro de Custo", "Descrição", "KM Percorrida", "Participação (%)"])
    for celula in resumo[1]:
        celula.font = NEGRITO

    rateios = fechamento.rateios.select_related("centro_custo").order_by("-km_percorrida")
    for rateio in rateios:
        resumo.append(
            [
                rateio.centro_custo.codigo,
                rateio.centro_custo.descricao,
                rateio.km_percorrida,
                float(rateio.percentual),
            ]
        )
    resumo.append(["Total", "", fechamento.total_km, 100.0])
    for celula in resumo[resumo.max_row]:
        celula.font = NEGRITO
    for linha in resumo.iter_rows(min_row=2, min_col=4, max_col=4):
        linha[0].number_format = "0.00"
    _autofit(resumo)

    viagens_sheet = wb.create_sheet("Viagens")
    viagens_sheet.append(
        ["Data", "Colaborador", "Centro de Custo", "KM Inicial", "KM Final", "KM Percorrida"]
    )
    for celula in viagens_sheet[1]:
        celula.font = NEGRITO

    viagens = (
        fechamento.viagens.select_related("funcionario", "centro_custo")
        .order_by("data", "id")
    )
    for viagem in viagens:
        viagens_sheet.append(
            [
                viagem.data.strftime("%d/%m/%Y"),
                str(viagem.funcionario),
                viagem.centro_custo.codigo,
                viagem.km_inicial,
                viagem.km_final,
                viagem.km_percorrida,
            ]
        )
    _autofit(viagens_sheet)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"fechamento_{fechamento.data_inicio:%Y%m%d}_{fechamento.data_fim:%Y%m%d}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'
    return response
